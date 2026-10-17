import csv, io, json, os
import numpy as np
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment

D = os.path.dirname(os.path.abspath(__file__))
CLASSES = ["airplane", "automobile", "bird", "cat", "deer", "dog", "frog", "horse", "ship", "truck"]
x = np.load(os.path.join(D, "set.npz"))["images"]
x = x.transpose(0, 2, 3, 1)
if x.dtype != np.uint8:
    x = ((x + 1) * 127.5 if x.min() < 0 else (x * 255 if x.max() <= 1 else x)).clip(0, 255).astype(np.uint8)
key = json.load(open(os.path.join(D, "key.json")))
rows = list(csv.DictReader(open(os.path.join(D, "my_eval.csv"))))

wb = Workbook(); ws = wb.active; ws.title = "Evaluation"
hdr = ["Image", "ID", "My realism", "My class", "My source", "True model", "True source",
       "True class", "Source correct", "Class correct"]
ws.append(hdr)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
ws.column_dimensions["A"].width = 12
for col, w in zip("BCDEFGHIJ", [6, 11, 15, 12, 14, 12, 12, 13, 13]):
    ws.column_dimensions[col].width = w

stats = {}
for i, r in enumerate(rows):
    model = key["sources"][i]
    true_src = "real" if model == "Real images" else "generated"
    lab = key["labels"][i]
    true_cls = CLASSES[lab] if lab >= 0 else "n/a"
    src_ok = r["source"] == true_src
    cls_ok = (r["class"] == true_cls) if lab >= 0 else None
    n = i + 2
    ws.append(["", int(r["id"]), int(r["realism"]), r["class"], r["source"], model, true_src,
               true_cls, "YES" if src_ok else "NO", "n/a" if cls_ok is None else ("YES" if cls_ok else "NO")])
    ws.row_dimensions[n].height = 60
    for c in ws[n]:
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws[f"I{n}"].fill = green if src_ok else red
    if cls_ok is not None:
        ws[f"J{n}"].fill = green if cls_ok else red
    buf = io.BytesIO(); Image.fromarray(x[i]).resize((76, 76), Image.NEAREST).save(buf, "PNG"); buf.seek(0)
    img = XLImage(buf); ws.add_image(img, f"A{n}")
    s = stats.setdefault(model, {"n": 0, "src_ok": 0, "realism": 0, "cls_ok": 0, "cls_n": 0})
    s["n"] += 1; s["src_ok"] += src_ok; s["realism"] += int(r["realism"])
    if cls_ok is not None:
        s["cls_n"] += 1; s["cls_ok"] += cls_ok
ws.freeze_panes = "B2"

sm = wb.create_sheet("Summary")
sm.append(["Model", "Images", "Mean my realism", "Source accuracy", "Class accuracy (real only)"])
for c in sm[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for m, s in stats.items():
    sm.append([m, s["n"], round(s["realism"] / s["n"], 2), f'{100*s["src_ok"]/s["n"]:.1f}%',
               f'{100*s["cls_ok"]/s["cls_n"]:.1f}%' if s["cls_n"] else "n/a"])
tot = sum(s["src_ok"] for s in stats.values())
sm.append(["ALL", len(rows), "", f"{100*tot/len(rows):.1f}%", ""])
for col, w in zip("ABCDE", [16, 10, 18, 18, 24]):
    sm.column_dimensions[col].width = w
out = os.path.join(D, "blind_evaluation.xlsx")
wb.save(out); print("saved", out); print(stats)
